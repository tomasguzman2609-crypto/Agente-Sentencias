from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.comments import Comment
from config import CAMPOS

FUENTE = "Arial"
COLOR_HEADER = "1F2937"
COLOR_REVISAR = "FDE68A"   # ámbar: hay algo que revisar
COLOR_OK = "D1FAE5"        # verde suave: todo verificado

COLUMNAS = list(CAMPOS.keys()) + ["¿REQUIERE REVISIÓN?", "NOMBRE DEL DOCUMENTO"]


def _valor_plano(campo_clave: str, datos: dict) -> str:
    """Texto principal de la celda: el valor + la página entre paréntesis
    (si hay cita), para que la referencia quede directamente a la vista."""
    if campo_clave == "resumen_causa":
        return datos.get("valor") or ""

    if campo_clave == "voto_disidente":
        if not datos.get("existe"):
            return "No"
        base = f"Sí — {datos.get('contenido') or ''}"
        pagina = datos.get("pagina")
        return f"{base} (pág. {pagina})" if pagina else base

    valor = datos.get("valor")
    if valor is None:
        return "No encontrado en el documento"
    pagina = datos.get("pagina")
    return f"{valor} (pág. {pagina})" if pagina else valor


def _comentario_cita(campo_clave: str, datos: dict):
    """Comentario de celda con la cita textual completa, para verificar
    el dato contra el documento original sin necesidad de un PDF aparte."""
    if campo_clave == "resumen_causa":
        return None
    cita = datos.get("cita_textual")
    if not cita:
        return None
    estado = datos.get("estado", "")
    return f"[{estado}] \u201c{cita}\u201d"


def _requiere_revision(extraccion_verificada: dict) -> bool:
    return any(
        datos.get("estado") in ("NO_VERIFICADO", "SINTESIS")
        for datos in extraccion_verificada.values()
    )


def construir_tabla(filas: list[dict], ruta_salida: str):
    """
    filas: lista de dicts con keys:
      - "extraccion": dict verificado (clave interna -> {"valor":..., "estado":...})
      - "nombre_archivo": str

    Cada celda muestra "valor (pág. N)" y, además, lleva un comentario de
    Excel con la cita textual completa que respalda ese valor — así la
    referencia queda directamente en la tabla, sin depender de un PDF aparte.
    """
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "Sentencias consolidadas"

    ws.append(COLUMNAS)
    for cell in ws[1]:
        cell.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=COLOR_HEADER)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    col_revisar = len(CAMPOS) + 1  # posición de la columna "¿REQUIERE REVISIÓN?"

    for fila in filas:
        extraccion = fila["extraccion"]
        claves = list(CAMPOS.values())
        valores = [_valor_plano(clave, extraccion[clave]) for clave in claves]
        revisar = "Sí" if _requiere_revision(extraccion) else "No"
        valores += [revisar, fila.get("nombre_archivo", "")]
        ws.append(valores)

        fila_actual = ws.max_row
        color = COLOR_REVISAR if revisar == "Sí" else COLOR_OK
        for cell in ws[fila_actual]:
            cell.font = Font(name=FUENTE, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=fila_actual, column=col_revisar).fill = PatternFill("solid", start_color=color)

        # Comentario con la cita textual completa, campo por campo
        for col_idx, clave in enumerate(claves, start=1):
            comentario = _comentario_cita(clave, extraccion[clave])
            if comentario:
                ws.cell(row=fila_actual, column=col_idx).comment = Comment(comentario, "Agente condenas")

    anchos = [18, 14, 14, 32, 22, 18, 24, 20, 24, 20, 26, 24, 16, 28]
    for i, ancho in enumerate(anchos[: len(COLUMNAS)], start=1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = ancho
    ws.freeze_panes = "A2"

    wb.save(ruta_salida)